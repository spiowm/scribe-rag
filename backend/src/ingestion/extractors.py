import re
import io
import threading
from google import genai

import pypdfium2
from pptx import Presentation
from google.genai import types


_PDFIUM_LOCK = threading.Lock()
_BASE64_IMG = re.compile(r"data:[^;]+;base64,[A-Za-z0-9+/=]+")


def pdf_text(blob: bytes) -> tuple[str, int]:
    with _PDFIUM_LOCK:
        doc = pypdfium2.PdfDocument(blob)
        try:
            pages = len(doc)
            parts = []
            for page in doc:
                textpage = page.get_textpage()
                try:
                    parts.append(textpage.get_text_bounded())
                finally:
                    textpage.close()
                    page.close()
            text = "\n".join(parts)
        finally:
            doc.close()
    return text, pages


def markdown_text(blob: bytes) -> tuple[str, int]:
    """Google Doc, експортований у text/markdown."""
    text = blob.decode("utf-8", errors="replace")
    text = _BASE64_IMG.sub("", text)
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)
    return text.strip(), 0


def pptx_text(blob: bytes) -> tuple[str, int]:
    """Текст і кількість слайдів."""
    prs = Presentation(io.BytesIO(blob))
    slides_text: list[str] = []

    for i, slide in enumerate(prs.slides, 1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                tf = getattr(shape, "text_frame", None)
                if tf and tf.text.strip():
                    parts.append(tf.text.strip())
            elif shape.has_table:
                table = getattr(shape, "table", None)
                if table:
                    for row in table.rows:
                        row_str = " | ".join(cell.text.strip() for cell in row.cells)
                        if row_str.strip():
                            parts.append(row_str)

        if slide.has_notes_slide:
            nts = getattr(slide.notes_slide.notes_text_frame, "text", None)
            if nts and nts.strip():
                parts.append(nts.strip())

        if parts:
            slides_text.append(f"[слайд {i}] " + "\n".join(parts))

    return "\n\n".join(slides_text), len(prs.slides)


EXTRACTION_PROMPT = """
Витягни з цього документа увесь текст: заголовки, абзаци, списки, підписи, текст на схемах і в таблицях. 
Зберігай порядок і структуру, таблиці подавай рядками через |. 
Не переказуй, не коментуй, не додавай вступів — лише текст документа."""


async def gemini_text(
    client: genai.Client,
    model: str,
    blob: bytes,
    mime: str = "application/pdf",
) -> str:

    resp = await client.aio.models.generate_content(
        model=model,
        contents=[types.Part.from_bytes(data=blob, mime_type=mime), EXTRACTION_PROMPT],
        config=types.GenerateContentConfig(temperature=0.0),
    )
    candidate = resp.candidates[0] if resp.candidates else None
    if candidate is None or not resp.text:
        raise ValueError(
            f"gemini не віддав текст: finish_reason={candidate.finish_reason if candidate else None}"
        )
    return resp.text
